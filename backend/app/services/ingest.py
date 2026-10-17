"""Turn uploaded files into validated deal and activity records.

Pure functions only: nothing here touches the database, so it is easy to test.

Supported inputs
* CRM export (CSV / XLSX): one row per deal. Column names are mapped through synonyms,
  so "Opportunity ID", "Amount" or "Close Date" all work.
* Activity log (CSV / XLSX): one row per email, call, meeting, note or follow-up.
* Documents (TXT / MD / DOCX / PDF / EML): email threads, call transcripts, meeting notes.
  Optional header lines at the top of each block ("Deal: DL-1042", "Date: 2026-09-16",
  "From: …", "Type: call") make matching exact; otherwise the deal is found from the
  account or deal name mentioned in the text. Separate several records in one file with
  a line of three or more dashes (---).
"""
from __future__ import annotations

import email
import email.policy
import io
import re
from dataclasses import dataclass, field
from typing import Any

import pandas as pd

from .utils import clean_str, is_blank, norm_header, parse_date, parse_money, parse_percent

TABLE_EXT = {".csv", ".xlsx", ".xls"}
DOC_EXT = {".txt", ".md", ".docx", ".pdf", ".eml"}
ALLOWED_EXT = TABLE_EXT | DOC_EXT

DEAL_COLUMNS: dict[str, list[str]] = {
    "deal_id": ["deal_id", "id", "opportunity_id", "opp_id", "record_id", "deal_number", "crm_id"],
    "name": ["deal_name", "name", "opportunity", "opportunity_name", "deal", "title"],
    "account": ["account", "account_name", "company", "customer", "organization", "organisation", "client"],
    "value": ["value", "amount", "deal_value", "arr", "acv", "deal_amount", "opportunity_amount", "deal_size"],
    "currency": ["currency", "currency_code"],
    "stage": ["stage", "sales_stage", "deal_stage", "pipeline_stage", "opportunity_stage"],
    "owner": ["owner", "deal_owner", "rep", "account_executive", "ae", "sales_rep", "opportunity_owner"],
    "expected_close": ["expected_close", "close_date", "expected_close_date", "closing_date", "target_close"],
    "last_activity": ["last_activity", "last_activity_date", "last_contacted", "last_touch", "last_contact"],
    "probability": ["probability", "win_probability", "close_probability", "likelihood", "probability_pct"],
    "forecast_category": ["forecast", "forecast_category", "forecast_status", "commit_status"],
    "status": ["status", "opportunity_status", "deal_status", "state", "opportunity_status"],
    "stage_entered": ["stage_entered", "stage_entered_date", "stage_date", "stage_since", "stage_changed"],
    "days_in_stage": ["days_in_stage", "stage_age", "stage_age_days"],
    "contacts": ["contacts", "stakeholders", "contact_list", "buying_committee"],
    "created": ["created", "created_date", "created_at", "first_contact", "opened"],
}

ACTIVITY_COLUMNS: dict[str, list[str]] = {
    "activity_id": ["activity_id", "record_id", "id", "message_id"],
    "deal_id": ["deal_id", "opportunity_id", "opp_id", "deal"],
    "account": ["account", "account_name", "company", "customer"],
    "type": ["type", "activity_type", "channel", "kind", "category"],
    "date": ["date", "timestamp", "activity_date", "sent_at", "occurred_at", "datetime", "sent", "created_at"],
    "direction": ["direction", "inbound_outbound", "in_out"],
    "from_": ["from", "sender", "from_email", "author", "from_name"],
    "to": ["to", "recipient", "recipients", "to_email"],
    "participants": ["participants", "attendees"],
    "subject": ["subject", "title", "topic"],
    "body": ["body", "content", "notes", "summary", "description", "text", "transcript", "message"],
    "outcome": ["outcome", "status", "result", "disposition"],
    "due_date": ["due_date", "due", "follow_up_date"],
}

TYPE_MAP = {
    "email": "email", "e_mail": "email", "mail": "email", "inbound_email": "email", "outbound_email": "email",
    "call": "call", "phone": "call", "phone_call": "call", "transcript": "call", "call_transcript": "call",
    "meeting": "meeting", "demo": "meeting", "zoom": "meeting", "video_call": "meeting", "meeting_notes": "meeting",
    "note": "note", "notes": "note", "internal_note": "note", "comment": "note",
    "follow_up": "follow_up", "followup": "follow_up", "task": "follow_up", "reminder": "follow_up", "todo": "follow_up",
}
TYPE_PREFIX = {"email": "EML", "call": "CALL", "meeting": "MTG", "note": "NOTE", "follow_up": "FUP"}

HEADER_KEYS = {
    "deal": "deal_id", "deal_id": "deal_id", "deal id": "deal_id", "opportunity": "deal_id", "opportunity id": "deal_id",
    "account": "account", "company": "account", "customer": "account",
    "type": "type", "activity": "type", "channel": "type",
    "date": "date", "sent": "date", "when": "date",
    "from": "from_", "sender": "from_", "to": "to", "cc": "cc",
    "subject": "subject", "title": "subject", "re": "subject",
    "direction": "direction", "outcome": "outcome", "status": "outcome", "result": "outcome",
    "participants": "participants", "attendees": "participants", "due": "due_date", "due date": "due_date",
}

ROLE_HINT = re.compile(r"\(([^)]+)\)")
EMAIL_HINT = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")


@dataclass
class ParseResult:
    kind: str
    records: list[dict[str, Any]] = field(default_factory=list)
    errors: list[dict[str, Any]] = field(default_factory=list)
    rows_total: int = 0


class UploadError(ValueError):
    """Raised for files that can't be read at all."""


# ---------------------------------------------------------------- tables

def extension(filename: str) -> str:
    m = re.search(r"(\.[A-Za-z0-9]+)$", filename or "")
    return m.group(1).lower() if m else ""


def read_tables(content: bytes, filename: str) -> list[tuple[str, pd.DataFrame]]:
    ext = extension(filename)
    if ext == ".csv":
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                df = pd.read_csv(io.BytesIO(content), dtype=str, keep_default_na=False, encoding=enc)
                return [("csv", df)]
            except UnicodeDecodeError:
                continue
            except pd.errors.EmptyDataError as exc:
                raise UploadError("The file is empty") from exc
            except pd.errors.ParserError as exc:
                raise UploadError(f"Could not read the CSV: {exc}") from exc
        raise UploadError("Could not decode the CSV file. Save it as UTF-8 and try again.")
    if ext in {".xlsx", ".xls"}:
        try:
            sheets = pd.read_excel(io.BytesIO(content), sheet_name=None, dtype=str, keep_default_na=False)
        except Exception as exc:  # openpyxl raises many types
            raise UploadError(f"Could not read the Excel file: {exc}") from exc
        out = [(name, df) for name, df in sheets.items() if not df.empty]
        if not out:
            raise UploadError("The workbook has no data")
        return out
    raise UploadError(f"{ext or 'This file type'} is not a table")


def map_columns(columns: list[Any], synonyms: dict[str, list[str]]) -> dict[str, str]:
    normalized = {norm_header(c): c for c in columns}
    mapping: dict[str, str] = {}
    used: set[str] = set()
    for canonical, names in synonyms.items():
        for n in names:
            if n in normalized and normalized[n] not in used:
                mapping[canonical] = normalized[n]
                used.add(normalized[n])
                break
    return mapping


def detect_kind(columns: list[Any]) -> str | None:
    d = map_columns(columns, DEAL_COLUMNS)
    a = map_columns(columns, ACTIVITY_COLUMNS)
    deal_score = sum(k in d for k in ("value", "stage", "expected_close", "probability", "owner"))
    act_score = sum(k in a for k in ("body", "type", "date", "subject", "from_"))
    if deal_score >= 2 and deal_score >= act_score:
        return "deals"
    if act_score >= 2 and "date" in a:
        return "activities"
    return None


def parse_contacts(raw: str | None) -> list[dict[str, str | None]]:
    """'Ana Ruiz (VP Operations) <ana@x.com>; Marcus Hale - IT Director - marcus@x.com'."""
    if not raw:
        return []
    contacts = []
    for part in re.split(r"[;|\n]", raw):
        part = part.strip()
        if not part:
            continue
        mail = EMAIL_HINT.search(part)
        email_addr = mail.group(0).lower() if mail else None
        role = None
        m = ROLE_HINT.search(part)
        if m:
            role = m.group(1).strip()
        rest = EMAIL_HINT.sub("", ROLE_HINT.sub("", part)).replace("<>", "").strip(" -,<>")
        bits = [b.strip() for b in re.split(r"\s+-\s+|,", rest) if b.strip()]
        name = bits[0] if bits else (email_addr.split("@")[0] if email_addr else None)
        if not role and len(bits) > 1:
            role = bits[1]
        if name:
            contacts.append({"name": name[:80], "role": (role or "")[:80] or None, "email": email_addr})
    return contacts


def _slug(*parts: str | None) -> str:
    s = "-".join(p for p in parts if p)
    return re.sub(r"[^A-Za-z0-9]+", "-", s).strip("-").upper()[:40] or "DEAL"


def parse_deal_rows(df: pd.DataFrame, sheet: str = "") -> ParseResult:
    res = ParseResult(kind="deals", rows_total=len(df))
    cols = map_columns(list(df.columns), DEAL_COLUMNS)
    missing = [c for c in ("value", "stage") if c not in cols]
    if "deal_id" not in cols and "name" not in cols:
        missing.append("deal_id or deal_name")
    if missing:
        res.errors.append({"row": None, "field": ", ".join(missing),
                           "message": f"Missing required column(s): {', '.join(missing)}"})
        return res
    seen: dict[str, int] = {}
    for i, row in df.iterrows():
        rownum = int(i) + 2  # header is row 1
        get = lambda k: row[cols[k]] if k in cols else None  # noqa: E731
        if all(is_blank(v) for v in row.values):
            res.rows_total -= 1
            continue
        errs = []
        name = clean_str(get("name"), 200)
        account = clean_str(get("account"), 200)
        deal_id = clean_str(get("deal_id"), 60) or _slug(account, name)
        value = parse_money(get("value"))
        if value is None or value < 0:
            errs.append(("value", f"'{get('value')}' is not a valid amount"))
        stage = clean_str(get("stage"), 60)
        if not stage:
            errs.append(("stage", "Stage is required"))
        exp_raw = get("expected_close")
        expected_close = parse_date(exp_raw)
        if not is_blank(exp_raw) and expected_close is None:
            errs.append(("expected_close", f"'{exp_raw}' is not a date"))
        prob_raw = get("probability")
        probability = parse_percent(prob_raw)
        if not is_blank(prob_raw) and probability is None:
            errs.append(("probability", f"'{prob_raw}' must be between 0 and 100"))
        if errs:
            for f, m in errs:
                res.errors.append({"row": rownum, "field": f, "message": m, "sheet": sheet or None})
            continue
        stage_entered = parse_date(get("stage_entered"))
        days_in_stage = None
        dis = get("days_in_stage")
        if not is_blank(dis):
            try:
                days_in_stage = int(float(str(dis)))
            except ValueError:
                pass
        status_raw = (clean_str(get("status")) or "open").lower()
        status = "won" if "won" in status_raw else "lost" if "lost" in status_raw else "open"
        rec = {
            "deal_id": deal_id,
            "name": name or deal_id,
            "account": account or name or deal_id,
            "value": float(value),
            "currency": (clean_str(get("currency"), 8) or "USD").upper(),
            "stage": stage,
            "owner": clean_str(get("owner"), 80) or "Unassigned",
            "expected_close": expected_close,
            "last_activity_date": parse_date(get("last_activity")),
            "probability": probability,
            "forecast_category": clean_str(get("forecast_category"), 40),
            "status": status,
            "stage_entered": stage_entered,
            "days_in_stage_reported": days_in_stage,
            "contacts": parse_contacts(clean_str(get("contacts"), 2000)),
            "crm_created": parse_date(get("created")),
        }
        if deal_id in seen:
            res.errors.append({"row": rownum, "field": "deal_id",
                               "message": f"Duplicate deal_id {deal_id} (row {seen[deal_id]}); the later row was used",
                               "level": "warning"})
            res.records = [r for r in res.records if r["deal_id"] != deal_id]
        seen[deal_id] = rownum
        res.records.append(rec)
    return res


# ---------------------------------------------------------------- matching

def deal_index(deals: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [{"deal_id": d["deal_id"], "account": d.get("account") or "", "name": d.get("name") or "",
             "owner": d.get("owner") or "", "contacts": d.get("contacts") or []} for d in deals]


def match_deal(index: list[dict[str, Any]], deal_ref: str | None = None, account: str | None = None,
               text: str = "") -> str | None:
    if deal_ref:
        ref = deal_ref.strip().lower()
        for d in index:
            if d["deal_id"].lower() == ref:
                return d["deal_id"]
        for d in index:
            if d["name"].lower() == ref or d["account"].lower() == ref:
                return d["deal_id"]
    if account:
        acc = account.strip().lower()
        hits = [d for d in index if d["account"].lower() == acc]
        if len(hits) == 1:
            return hits[0]["deal_id"]
        hits = [d for d in index if acc and (acc in d["account"].lower() or d["account"].lower() in acc)]
        if len(hits) == 1:
            return hits[0]["deal_id"]
    if text:
        low = text.lower()
        counts: dict[str, int] = {}
        for d in index:
            n = 0
            if d["deal_id"] and re.search(rf"\b{re.escape(d['deal_id'].lower())}\b", low):
                n += 5
            for term in (d["account"], d["name"]):
                if term and len(term) >= 4:
                    n += low.count(term.lower())
            for c in d["contacts"]:
                if c.get("email") and c["email"] in low:
                    n += 2
            if n:
                counts[d["deal_id"]] = n
        if counts:
            ranked = sorted(counts.items(), key=lambda kv: kv[1], reverse=True)
            if len(ranked) == 1 or ranked[0][1] > ranked[1][1]:
                return ranked[0][0]
    return None


def normalize_type(raw: str | None, filename: str = "", text: str = "") -> str:
    if raw:
        key = norm_header(raw)
        if key in TYPE_MAP:
            return TYPE_MAP[key]
        for k, v in TYPE_MAP.items():
            if k in key:
                return v
    hint = (filename or "").lower()
    if any(w in hint for w in ("transcript", "call")):
        return "call"
    if any(w in hint for w in ("meeting", "minutes")):
        return "meeting"
    if any(w in hint for w in ("follow", "task")):
        return "follow_up"
    if any(w in hint for w in ("email", "mail", "thread", ".eml")):
        return "email"
    if text and re.search(r"^\s*(from|to|subject):", text, re.I | re.M):
        return "email"
    return "note"


def _person_match(who: str, person_name: str | None, person_email: str | None) -> bool:
    w = who.lower()
    if person_email and person_email.lower() in w:
        return True
    if person_name:
        parts = [p for p in re.findall(r"[a-z]+", person_name.lower()) if len(p) > 1]
        if parts and all(p in w for p in parts):
            return True
        if parts and parts[0] in w and len(parts[0]) >= 4:
            return True
    return False


def infer_direction(rec: dict[str, Any], deal: dict[str, Any] | None, explicit: str | None = None) -> str:
    if explicit:
        e = explicit.strip().lower()
        if e.startswith(("in", "rec", "cust")):
            return "inbound"
        if e.startswith(("out", "sent", "rep")):
            return "outbound"
        if e.startswith(("int",)):
            return "internal"
    t = rec.get("type")
    if t in ("call", "meeting"):
        return "two_way"
    if t in ("note", "follow_up"):
        return "internal"
    sender = rec.get("from_") or ""
    if deal and sender:
        if _person_match(sender, deal.get("owner"), None):
            return "outbound"
        for c in deal.get("contacts") or []:
            if _person_match(sender, c.get("name"), c.get("email")):
                return "inbound"
        recipient = rec.get("to") or ""
        for c in deal.get("contacts") or []:
            if recipient and _person_match(recipient, c.get("name"), c.get("email")):
                return "outbound"
    return "inbound" if t == "email" and sender else "unknown"


def _outcome(raw: str | None) -> str | None:
    if not raw:
        return None
    r = raw.strip().lower()
    for key, val in (("declin", "declined"), ("cancel", "cancelled"), ("no show", "no_show"), ("no-show", "no_show"),
                     ("noshow", "no_show"), ("resched", "rescheduled"), ("miss", "missed"), ("overdue", "missed"),
                     ("done", "completed"), ("complet", "completed"), ("held", "completed"), ("attend", "completed"),
                     ("pending", "pending"), ("open", "pending"), ("scheduled", "scheduled"), ("no answer", "no_answer"),
                     ("voicemail", "no_answer")):
        if key in r:
            return val
    return r[:40]


def build_activity(*, deal_id: str, type_: str, date, direction: str, from_: str | None, to: str | None,
                   subject: str | None, body: str | None, outcome: str | None, participants: str | None,
                   due_date, source_file: str, activity_id: str | None = None) -> dict[str, Any]:
    return {
        "deal_id": deal_id, "activity_id": activity_id, "type": type_, "date": date, "direction": direction,
        "from_": from_, "to": to, "subject": subject, "body": body, "outcome": outcome,
        "participants": participants, "due_date": due_date, "source_file": source_file,
    }


def parse_activity_rows(df: pd.DataFrame, deals: list[dict[str, Any]], filename: str, sheet: str = "") -> ParseResult:
    res = ParseResult(kind="activities", rows_total=len(df))
    cols = map_columns(list(df.columns), ACTIVITY_COLUMNS)
    if "date" not in cols:
        res.errors.append({"row": None, "field": "date", "message": "Missing required column: date"})
        return res
    if "body" not in cols and "subject" not in cols:
        res.errors.append({"row": None, "field": "body", "message": "Need a body/notes or subject column"})
        return res
    index = deal_index(deals)
    by_id = {d["deal_id"]: d for d in deals}
    for i, row in df.iterrows():
        rownum = int(i) + 2
        if all(is_blank(v) for v in row.values):
            res.rows_total -= 1
            continue
        get = lambda k: clean_str(row[cols[k]]) if k in cols else None  # noqa: E731
        body, subject = get("body"), get("subject")
        dt = parse_date(row[cols["date"]])
        if dt is None:
            res.errors.append({"row": rownum, "field": "date", "message": f"'{row[cols['date']]}' is not a date"})
            continue
        if not body and not subject:
            res.errors.append({"row": rownum, "field": "body", "message": "Row has no content"})
            continue
        deal_id = match_deal(index, get("deal_id"), get("account"), f"{subject or ''}\n{body or ''}")
        if not deal_id:
            res.errors.append({"row": rownum, "field": "deal_id",
                               "message": "Could not match this row to a deal. Add a deal_id or account column."})
            continue
        type_ = normalize_type(get("type"), filename)
        rec = build_activity(
            deal_id=deal_id, type_=type_, date=dt, direction="unknown", from_=get("from_"), to=get("to"),
            subject=subject, body=body, outcome=_outcome(get("outcome")), participants=get("participants"),
            due_date=parse_date(get("due_date")), source_file=filename, activity_id=get("activity_id"),
        )
        rec["direction"] = infer_direction(rec, by_id.get(deal_id), get("direction"))
        res.records.append(rec)
    return res


# ---------------------------------------------------------------- documents

def extract_text(content: bytes, filename: str) -> str:
    ext = extension(filename)
    if ext in {".txt", ".md"}:
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                return content.decode(enc)
            except UnicodeDecodeError:
                continue
        raise UploadError("Could not decode the text file")
    if ext == ".docx":
        try:
            from docx import Document
        except ImportError as exc:  # pragma: no cover
            raise UploadError("python-docx is not installed") from exc
        try:
            doc = Document(io.BytesIO(content))
        except Exception as exc:
            raise UploadError(f"Could not read the Word file: {exc}") from exc
        parts = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for r in t.rows:
                parts.append(" | ".join(c.text for c in r.cells))
        return "\n".join(parts)
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:  # pragma: no cover
            raise UploadError("pypdf is not installed") from exc
        try:
            reader = PdfReader(io.BytesIO(content))
            text = "\n".join((p.extract_text() or "") for p in reader.pages)
        except Exception as exc:
            raise UploadError(f"Could not read the PDF: {exc}") from exc
        if not text.strip():
            raise UploadError("This PDF has no selectable text (it may be a scan)")
        return text
    if ext == ".eml":
        msg = email.message_from_bytes(content, policy=email.policy.default)
        body_part = msg.get_body(preferencelist=("plain", "html"))
        body = body_part.get_content() if body_part else ""
        if body_part is not None and body_part.get_content_type() == "text/html":
            body = re.sub(r"<[^>]+>", " ", body)
        headers = [f"{k}: {msg[k]}" for k in ("From", "To", "Subject", "Date") if msg[k]]
        return "\n".join(headers) + "\n\n" + body
    raise UploadError(f"{ext or 'This file type'} is not supported")


def split_blocks(text: str) -> list[str]:
    blocks = re.split(r"(?m)^\s*(?:-{3,}|={3,}|_{3,})\s*$", text)
    return [b.strip() for b in blocks if b.strip()]


def parse_headers(block: str) -> tuple[dict[str, str], str]:
    headers: dict[str, str] = {}
    lines = block.splitlines()
    body_start = 0
    for idx, line in enumerate(lines):
        if not line.strip():
            if headers:
                body_start = idx + 1
                break
            continue
        m = re.match(r"^\s*([A-Za-z][A-Za-z _]{0,20}):\s*(.*)$", line)
        key = HEADER_KEYS.get(m.group(1).strip().lower()) if m else None
        if not key:
            body_start = idx
            break
        headers.setdefault(key, m.group(2).strip())
        body_start = idx + 1
    body = "\n".join(lines[body_start:]).strip()
    return headers, body


def parse_document(content: bytes, filename: str, deals: list[dict[str, Any]]) -> ParseResult:
    res = ParseResult(kind="documents")
    text = extract_text(content, filename)
    index = deal_index(deals)
    by_id = {d["deal_id"]: d for d in deals}
    file_date = None
    m = re.search(r"(20\d{2})[-_](\d{2})[-_](\d{2})", filename)
    if m:
        file_date = parse_date(f"{m[1]}-{m[2]}-{m[3]}")
    blocks = split_blocks(text)
    res.rows_total = len(blocks)
    for n, block in enumerate(blocks, start=1):
        headers, body = parse_headers(block)
        dt = parse_date(headers.get("date")) or file_date
        if dt is None:
            res.errors.append({"row": n, "field": "date",
                               "message": "No date found. Add a line like 'Date: 2026-09-16' at the top of this section."})
            continue
        if not body and not headers.get("subject"):
            res.errors.append({"row": n, "field": "body", "message": "Section has no content"})
            continue
        deal_id = match_deal(index, headers.get("deal_id"), headers.get("account"), block)
        if not deal_id:
            res.errors.append({"row": n, "field": "deal_id",
                               "message": "Could not tell which deal this is about. Add 'Deal: <deal id>' at the top."})
            continue
        type_ = normalize_type(headers.get("type"), filename, block)
        rec = build_activity(
            deal_id=deal_id, type_=type_, date=dt, direction="unknown", from_=headers.get("from_"),
            to=headers.get("to"), subject=headers.get("subject"), body=body[:20000] or None,
            outcome=_outcome(headers.get("outcome")), participants=headers.get("participants"),
            due_date=parse_date(headers.get("due_date")), source_file=filename,
        )
        rec["direction"] = infer_direction(rec, by_id.get(deal_id), headers.get("direction"))
        res.records.append(rec)
    return res
